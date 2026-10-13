from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse, unquote
import json, os, threading, shutil, uuid, io
import publication_store as ledger
from datetime import datetime, timezone
ROOT=Path(__file__).resolve().parent
DATA=Path(os.environ.get('STUDIO_DATA',ROOT/'data/topics.json'))
PUBLICATIONS=Path(os.environ.get('STUDIO_PUBLICATIONS',DATA.parent/'publications.json'))
PORT=int(os.environ.get('STUDIO_PORT','19826'))
LOCK=threading.RLock()
STATES=['选题库','聊想法','写稿中','待拍','已拍']
def now():return datetime.now(timezone.utc).isoformat()
def read():return json.loads(DATA.read_text())
def save(data):
 data['updatedAt']=now();DATA.parent.mkdir(exist_ok=True,parents=True)
 backups=DATA.parent/'backups';backups.mkdir(exist_ok=True)
 if DATA.exists():shutil.copy2(DATA,backups/(datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'.json'))
 for p in sorted(backups.glob('*.json'))[:-30]:p.unlink()
 tmp=DATA.with_suffix('.tmp');tmp.write_text(json.dumps(data,ensure_ascii=False,indent=2));os.replace(tmp,DATA)
def validate_topic(t):
 if not isinstance(t,dict) or not isinstance(t.get('title'),str) or not t['title'].strip():raise ValueError('请填写选题名称')
 if t.get('status') not in STATES:raise ValueError('无效状态')
 if t.get('priority','待评估') not in ['优先拍','备选','暂缓','待评估']:raise ValueError('无效优先级')
 if 'priorityOrder' in t and (type(t['priorityOrder']) is not int or not 1<=t['priorityOrder']<=999):raise ValueError('无效排序')
 for key in ['category','original','hook','angle','notes','script','priorityReason','priorityNext','prioritySignal','priorityReviewedAt']:
  if not isinstance(t.get(key,''),str) or len(t.get(key,''))>100000:raise ValueError('内容格式错误')
 for key in ['evidence','questions','related']:
  if not isinstance(t.get(key,[]),list):raise ValueError('列表格式错误')
 for e in t.get('evidence',[]):
  if not isinstance(e,dict):raise ValueError('信源格式错误')
  if e.get('url') and urlparse(e['url']).scheme not in ['http','https']:raise ValueError('链接需以https或http开头')
def word(t):
 from docx import Document
 from docx.shared import Pt
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 d=Document();st=d.styles['Normal'];st.font.name='FangSong';st.font.size=Pt(12);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'FangSong')
 d.add_heading(t['title'],0)
 for label,key in [('口播稿','script'),('开场','hook'),('观点','angle'),('我的想法','notes')]:
  if t.get(key):d.add_heading(label,1);d.add_paragraph(t[key])
 if t.get('evidence'):
  d.add_heading('资料来源',1)
  for e in t['evidence']:
   d.add_paragraph(e['title']);d.add_paragraph(e.get('statement',''))
   if e.get('url'):
    p=d.add_paragraph();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(e['url'],RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');text=OxmlElement('w:t');text.text=e['url'];r.append(text);h.append(r);p._p.append(h)
 b=io.BytesIO();d.save(b);return b.getvalue()
class Handler(SimpleHTTPRequestHandler):
 def __init__(self,*a,**kw):super().__init__(*a,directory=str(ROOT/'dist'),**kw)
 def jsonout(self,x,status=200):
  b=json.dumps(x,ensure_ascii=False).encode();self.send_response(status);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Cache-Control','no-store');self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b)
 def allowed(self):
  host=self.headers.get('Host','').split(':')[0]
  origin=self.headers.get('Origin')
  return host in ['127.0.0.1','localhost'] and (not origin or urlparse(origin).netloc==self.headers.get('Host'))
 def do_GET(self):
  if not self.allowed():return self.jsonout({'error':'仅支持本机访问'},403)
  path=urlparse(self.path).path
  if path=='/api/publications':
   with LOCK:return self.jsonout(ledger.load(PUBLICATIONS))
  if path in ['/api/data','/api/backup']:
   with LOCK:
    d=read()
    if path=='/api/backup':d['publicationLedger']=ledger.load(PUBLICATIONS)
    return self.jsonout(d)
  if path.startswith('/api/word/'):
   with LOCK:t=next((t for t in read()['topics'] if t['id']==unquote(path.split('/')[-1])),None)
   if not t:return self.jsonout({'error':'选题不存在'},404)
   b=word(t);self.send_response(200);self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.wordprocessingml.document');self.send_header('Content-Disposition','attachment; filename="topic.docx"');self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b);return
  if path.startswith('/api/'):return self.jsonout({'error':'找不到接口'},404)
  if path=='/':self.path='/index.html'
  return super().do_GET()
 def do_POST(self):
  if not self.allowed():return self.jsonout({'error':'仅支持本机访问'},403)
  try:
   size=int(self.headers.get('Content-Length','0'))
   if size>2000000:raise ValueError('内容过大')
   body=json.loads(self.rfile.read(size))
   with LOCK:
    d=read();path=urlparse(self.path).path
    if path in ['/api/publications/save','/api/publications/import']:
     records=ledger.load(PUBLICATIONS)
     if path.endswith('/save') and body.get('revision')!=records.get('revision',0):return self.jsonout({'error':'发布台账已更新，请刷新后保存。'},409)
     records=ledger.merge(records,body)
     ledger.validate(records,{t['id'] for t in d['topics']})
     ledger.save(PUBLICATIONS,records);return self.jsonout(records)
    if path=='/api/topic':
     t=body['topic'];validate_topic(t)
     existing=next((x for x in d['topics'] if x['id']==t.get('id')),None)
     if existing:
      if body.get('revision')!=existing.get('revision',0):return self.jsonout({'error':'这条选题已在另一窗口更新，请刷新后再保存。'},409)
      t['history']=(existing.get('history',[])+[{'at':now(),'status':existing['status'],'notes':existing['notes'],'script':existing['script'],'hook':existing['hook'],'priority':existing.get('priority','待评估'),'priorityReason':existing.get('priorityReason',''),'priorityNext':existing.get('priorityNext','')}])[-20:]
      t['revision']=existing.get('revision',0)+1;d['topics'][d['topics'].index(existing)]=t
     else:
      t['id']=str(uuid.uuid4());t['revision']=0;t['history']=[];d['topics'].insert(0,t)
     t['updatedAt']=now();save(d);return self.jsonout({'topic':t})
    if path=='/api/preferences':
     prefs=body.get('preferences')
     if not isinstance(prefs,dict) or any(not isinstance(x,str) for x in prefs.values()):raise ValueError('偏好格式错误')
     d['preferences']=prefs;save(d);return self.jsonout(d['preferences'])
    return self.jsonout({'error':'找不到接口'},404)
  except (ValueError,KeyError,TypeError) as e:return self.jsonout({'error':str(e)},400)
  except Exception as e:print(e,flush=True);return self.jsonout({'error':'保存失败，原有数据仍保留。'},500)
 def log_message(self,fmt,*args):print(fmt%args,flush=True)
if __name__=='__main__':
 print(f'创作工作台 http://127.0.0.1:{PORT}',flush=True)
 ThreadingHTTPServer(('127.0.0.1',PORT),Handler).serve_forever()
